"""Export service for PDF, Markdown, JSON, and DOCX manuscripts, plus file import helpers."""

import io
import json
import logging
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple

from storybook.config import config
from storybook.models import Project, Chapter, Scene, Character, Location, CanonFact, StoryEvent
from storybook.repositories import (
    ProjectRepository,
    ChapterRepository,
    SceneRepository,
    StoryBibleRepository,
)

logger = logging.getLogger(__name__)


class ExportService:
    """Manages document rendering, exports (PDF, Markdown, JSON, DOCX), and manuscript imports."""

    def __init__(
        self,
        project_repo: Optional[ProjectRepository] = None,
        chapter_repo: Optional[ChapterRepository] = None,
        scene_repo: Optional[SceneRepository] = None,
        bible_repo: Optional[StoryBibleRepository] = None,
    ):
        self.project_repo = project_repo or ProjectRepository()
        self.chapter_repo = chapter_repo or ChapterRepository()
        self.scene_repo = scene_repo or SceneRepository()
        self.bible_repo = bible_repo or StoryBibleRepository()

    # Markdown Export
    def export_markdown(self, project_id: str) -> str:
        """Export project as a clean, complete Markdown manuscript."""
        project = self.project_repo.get(project_id)
        if not project:
            raise ValueError(f"Project '{project_id}' not found.")

        chapters = self.chapter_repo.list_by_project(project_id)
        lines = [
            f"# {project.title}",
            f"*{project.subtitle}*" if project.subtitle else "",
            "",
            f"**Genre:** {project.genre} | **Tone:** {project.tone} | **Language:** {project.language}",
            "",
            "## Synopsis",
            project.synopsis or project.premise or "No synopsis available.",
            "",
            "---",
            "",
            "## Table of Contents",
        ]

        for ch in chapters:
            lines.append(f"- [{ch.title}](#{ch.title.lower().replace(' ', '-')})")

        lines.extend(["", "---", ""])

        for ch in chapters:
            lines.append(f"## {ch.title}")
            if ch.summary:
                lines.append(f"> *{ch.summary}*\n")

            scenes = self.scene_repo.list_by_chapter(ch.id)
            if scenes:
                for sc in scenes:
                    lines.append(f"### {sc.title}")
                    if sc.prose:
                        lines.append(sc.prose.strip())
                    elif sc.summary:
                        lines.append(f"*[Scene Outline: {sc.summary}]*")
                    lines.append("")
            elif ch.prose:
                lines.append(ch.prose.strip())
                lines.append("")

            lines.append("---")
            lines.append("")

        return "\n".join(lines)

    # JSON Export (Future Image-Pipeline Compatible)
    def export_json(self, project_id: str) -> Dict[str, Any]:
        """Export full project hierarchy as JSON with complete future image schema."""
        project = self.project_repo.get(project_id)
        if not project:
            raise ValueError(f"Project '{project_id}' not found.")

        chapters = self.chapter_repo.list_by_project(project_id)
        characters = self.bible_repo.list_characters(project_id)
        locations = self.bible_repo.list_locations(project_id)
        canon_facts = self.bible_repo.list_canon_facts(project_id)
        events = self.bible_repo.list_story_events(project_id)

        chapters_payload = []
        for ch in chapters:
            scenes = self.scene_repo.list_by_chapter(ch.id)
            scenes_payload = []
            for sc in scenes:
                scenes_payload.append({
                    "id": sc.id,
                    "order_index": sc.order_index,
                    "title": sc.title,
                    "summary": sc.summary,
                    "prose": sc.prose,
                    "purpose": sc.purpose,
                    "pov": sc.pov,
                    "location_name": sc.location_name,
                    "time_marker": sc.time_marker,
                    "characters_present": sc.characters_present,
                    "dramatic_conflict": sc.continuity_json.get("dramatic_conflict", ""),
                    "emotional_shift": sc.emotional_shift,
                    "continuity": sc.continuity_json,
                    "status": sc.status,
                    # Image & Comic Panel Generation Schema (Phase 1 Ready)
                    "visual_metadata": {
                        "camera_framing": sc.camera_framing or "Standard shot",
                        "mood_atmosphere": sc.mood_atmosphere or "Neutral narrative mood",
                        "visual_summary": sc.visual_summary or sc.summary,
                        "character_visual_refs": sc.characters_present,
                    }
                })

            chapters_payload.append({
                "id": ch.id,
                "order_index": ch.order_index,
                "title": ch.title,
                "summary": ch.summary,
                "prose": ch.prose,
                "status": ch.status,
                "scenes": scenes_payload,
            })

        return {
            "schema_version": "1.0.0",
            "project": {
                "id": project.id,
                "title": project.title,
                "subtitle": project.subtitle,
                "genre": project.genre,
                "tone": project.tone,
                "language": project.language,
                "intended_audience": project.intended_audience,
                "premise": project.premise,
                "synopsis": project.synopsis,
                "created_at": project.created_at,
                "updated_at": project.updated_at,
            },
            "story_bible": {
                "characters": [c.model_dump() for c in characters],
                "locations": [l.model_dump() for l in locations],
                "canon_facts": [f.model_dump() for f in canon_facts],
                "timeline_events": [e.model_dump() for e in events],
            },
            "chapters": chapters_payload,
        }

    # PDF Export via ReportLab
    def export_pdf(self, project_id: str) -> bytes:
        """Generate a professionally formatted book-like PDF using ReportLab."""
        project = self.project_repo.get(project_id)
        if not project:
            raise ValueError(f"Project '{project_id}' not found.")

        chapters = self.chapter_repo.list_by_project(project_id)

        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.platypus import (
                SimpleDocTemplate,
                Paragraph,
                Spacer,
                PageBreak,
                HRFlowable,
            )
            from reportlab.pdfgen import canvas

            class NumberedCanvas(canvas.Canvas):
                """Two-pass canvas to calculate and print total page numbers."""
                def __init__(self, *args, **kwargs):
                    super().__init__(*args, **kwargs)
                    self._saved_page_states = []

                def showPage(self):
                    self._saved_page_states.append(dict(self.__dict__))
                    self._startPage()

                def save(self):
                    num_pages = len(self._saved_page_states)
                    for state in self._saved_page_states:
                        self.__dict__.update(state)
                        self.draw_page_number(num_pages)
                        super().showPage()
                    super().save()

                def draw_page_number(self, page_count):
                    if self._pageNumber > 1:  # Omit page number on cover page
                        self.setFont("Helvetica", 9)
                        text = f"Page {self._pageNumber} of {page_count}"
                        self.drawRightString(letter[0] - 54, 36, text)

            buffer = io.BytesIO()
            doc = SimpleDocTemplate(
                buffer,
                pagesize=letter,
                leftMargin=54,
                rightMargin=54,
                topMargin=54,
                bottomMargin=54,
            )

            styles = getSampleStyleSheet()

            # Custom typography styles
            title_style = ParagraphStyle(
                "CoverTitle",
                parent=styles["Title"],
                fontName="Helvetica-Bold",
                fontSize=28,
                leading=34,
                alignment=1,  # Centered
                spaceAfter=15,
            )
            subtitle_style = ParagraphStyle(
                "CoverSubtitle",
                parent=styles["Normal"],
                fontName="Helvetica-Oblique",
                fontSize=14,
                leading=18,
                alignment=1,
                spaceAfter=30,
            )
            synopsis_style = ParagraphStyle(
                "CoverSynopsis",
                parent=styles["Normal"],
                fontName="Helvetica",
                fontSize=11,
                leading=16,
                alignment=4,  # Justified
                spaceBefore=20,
            )
            chapter_title_style = ParagraphStyle(
                "ChapterTitle",
                parent=styles["Heading1"],
                fontName="Helvetica-Bold",
                fontSize=18,
                leading=22,
                spaceBefore=15,
                spaceAfter=10,
            )
            scene_title_style = ParagraphStyle(
                "SceneTitle",
                parent=styles["Heading2"],
                fontName="Helvetica-Bold",
                fontSize=13,
                leading=16,
                spaceBefore=12,
                spaceAfter=6,
            )
            body_style = ParagraphStyle(
                "BookProse",
                parent=styles["Normal"],
                fontName="Times-Roman",
                fontSize=11,
                leading=16,
                spaceAfter=10,
                alignment=4,  # Justified
                firstLineIndent=18,
            )
            summary_style = ParagraphStyle(
                "ChapterSummary",
                parent=styles["Italic"],
                fontName="Times-Italic",
                fontSize=10,
                leading=14,
                spaceAfter=12,
            )

            story = []

            # 1. Cover / Title Page
            story.append(Spacer(1, 2 * inch))
            story.append(Paragraph(project.title, title_style))
            if project.subtitle:
                story.append(Paragraph(project.subtitle, subtitle_style))
            story.append(Spacer(1, 0.5 * inch))

            meta_text = f"<b>Genre:</b> {project.genre} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Tone:</b> {project.tone}"
            story.append(Paragraph(meta_text, ParagraphStyle("Meta", alignment=1, fontSize=10)))

            if project.synopsis:
                story.append(Spacer(1, 0.8 * inch))
                story.append(Paragraph(f"<b>Synopsis:</b> {project.synopsis}", synopsis_style))

            story.append(PageBreak())

            # 2. Table of Contents
            story.append(Paragraph("Table of Contents", chapter_title_style))
            story.append(HRFlowable(width="100%", thickness=1, spaceAfter=15))
            for ch in chapters:
                toc_line = f"• <b>{ch.title}</b>"
                story.append(Paragraph(toc_line, styles["Normal"]))
                story.append(Spacer(1, 4))
            story.append(PageBreak())

            # 3. Chapters & Scenes
            for ch in chapters:
                story.append(Paragraph(ch.title, chapter_title_style))
                story.append(HRFlowable(width="100%", thickness=0.5, spaceAfter=10))

                if ch.summary:
                    story.append(Paragraph(f"<i>{ch.summary}</i>", summary_style))

                scenes = self.scene_repo.list_by_chapter(ch.id)
                if scenes:
                    for sc in scenes:
                        story.append(Paragraph(sc.title, scene_title_style))
                        prose_content = sc.prose or sc.summary or "Draft in progress."
                        # Split prose paragraphs
                        for para in prose_content.split("\n\n"):
                            clean_para = para.strip().replace("\n", " ")
                            if clean_para:
                                story.append(Paragraph(clean_para, body_style))
                        story.append(Spacer(1, 8))
                elif ch.prose:
                    for para in ch.prose.split("\n\n"):
                        clean_para = para.strip().replace("\n", " ")
                        if clean_para:
                            story.append(Paragraph(clean_para, body_style))

                story.append(PageBreak())

            doc.build(story, canvasmaker=NumberedCanvas)
            return buffer.getvalue()

        except Exception as e:
            logger.error(f"PDF generation failed: {e}")
            raise RuntimeError(f"ReportLab PDF generation error: {e}") from e

    # DOCX Export via python-docx
    def export_docx(self, project_id: str) -> bytes:
        """Export project as a formatted Microsoft Word (.docx) document."""
        project = self.project_repo.get(project_id)
        if not project:
            raise ValueError(f"Project '{project_id}' not found.")

        chapters = self.chapter_repo.list_by_project(project_id)

        try:
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document()

            # Title & Subtitle
            title_p = doc.add_paragraph()
            title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = title_p.add_run(project.title)
            run.font.size = Pt(26)
            run.font.bold = True

            if project.subtitle:
                sub_p = doc.add_paragraph()
                sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                s_run = sub_p.add_run(project.subtitle)
                s_run.font.size = Pt(14)
                s_run.font.italic = True

            meta_p = doc.add_paragraph()
            meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            m_run = meta_p.add_run(f"Genre: {project.genre} | Tone: {project.tone}")
            m_run.font.size = Pt(10)
            m_run.font.color.rgb = RGBColor(128, 128, 128)

            doc.add_page_break()

            # Chapters
            for ch in chapters:
                doc.add_heading(ch.title, level=1)
                if ch.summary:
                    sum_p = doc.add_paragraph()
                    sum_run = sum_p.add_run(f"Chapter Summary: {ch.summary}")
                    sum_run.italic = True

                scenes = self.scene_repo.list_by_chapter(ch.id)
                if scenes:
                    for sc in scenes:
                        doc.add_heading(sc.title, level=2)
                        prose_content = sc.prose or sc.summary or ""
                        for para in prose_content.split("\n\n"):
                            clean_para = para.strip().replace("\n", " ")
                            if clean_para:
                                doc.add_paragraph(clean_para)
                elif ch.prose:
                    for para in ch.prose.split("\n\n"):
                        clean_para = para.strip().replace("\n", " ")
                        if clean_para:
                            doc.add_paragraph(clean_para)

                doc.add_page_break()

            buffer = io.BytesIO()
            doc.save(buffer)
            return buffer.getvalue()
        except Exception as e:
            logger.error(f"DOCX generation failed: {e}")
            raise RuntimeError(f"DOCX export error: {e}") from e

    # Manuscript Import Helper
    @staticmethod
    def extract_text_from_file(filename: str, file_bytes: bytes) -> Tuple[str, str]:
        """
        Extract text from uploaded TXT, MD, PDF, or DOCX files.
        Returns (extracted_text, detected_format).
        """
        lower = filename.lower()
        if lower.endswith(".txt") or lower.endswith(".md"):
            try:
                return file_bytes.decode("utf-8"), "txt/md"
            except UnicodeDecodeError:
                return file_bytes.decode("latin-1"), "txt/md"

        elif lower.endswith(".pdf"):
            try:
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(file_bytes))
                pages = [page.extract_text() or "" for page in reader.pages]
                return "\n\n".join(pages).strip(), "pdf"
            except Exception as e:
                raise ValueError(f"Failed to extract text from PDF: {e}") from e

        elif lower.endswith(".docx"):
            try:
                from docx import Document
                doc = Document(io.BytesIO(file_bytes))
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                return "\n\n".join(paragraphs).strip(), "docx"
            except Exception as e:
                raise ValueError(f"Failed to extract text from DOCX: {e}") from e

        else:
            raise ValueError(f"Unsupported file format '{filename}'. Supported: .txt, .md, .pdf, .docx")


export_service = ExportService()
